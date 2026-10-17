"""自動統計分析報告：依資料型態與前提檢定挑選方法，產生 seaborn 圖表、統計表與中英文解讀，可輸出 HTML / Word / PPT / Excel。

文字以 i18n._() 樣板產生 (中文為原文)；內部表格鍵值維持中文，報告完成後統一以 i18n.table() 轉換欄名與文字。
"""

from __future__ import annotations

import base64
import datetime as dt
import html
import io
import math
import platform
import re
from dataclasses import dataclass, field

import matplotlib
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import scipy
import seaborn as sns
import statsmodels
from scipy import stats as sps

from . import __version__, data, export, fonts, i18n, plots
from . import stats as bs
from .i18n import _
from .style import Style, build_rc, tidy_xticklabels

ALPHA = bs.ALPHA
STRATEGIES = {"自動 (依前提檢定選擇)": "auto", "一律母數檢定": "parametric", "一律無母數檢定": "nonparametric"}

METHOD_LABEL = {
    "student": "Student t 檢定",
    "welch": "Welch t 檢定",
    "mwu": "Mann-Whitney U 檢定",
    "anova": "單因子 ANOVA + Tukey HSD 事後檢定",
    "welch_anova": "Welch ANOVA + Games-Howell 事後檢定",
    "kw": "Kruskal-Wallis 檢定 + Dunn 事後檢定 (Holm 校正)",
}
_TEST_ROW = {"student": "Student t 檢定 (等變異)", "welch": "Welch t 檢定 (不等變異)", "mwu": "Mann-Whitney U (無母數)",
             "anova": "One-way ANOVA", "welch_anova": "Welch ANOVA (不等變異)", "kw": "Kruskal-Wallis (無母數)"}
_POSTHOC = {"anova": "Tukey HSD", "welch_anova": "Games-Howell", "kw": "Dunn (Holm)"}
_SIG_TEST = {"student": "student", "welch": "welch", "mwu": "mwu", "anova": "tukey", "welch_anova": "gh", "kw": "dunn"}
_DESC_COLS = ["變項", "組別", "n", "缺失數", "平均值 Mean", "標準差 SD", "標準誤 SEM", "95% CI 下限", "95% CI 上限",
              "中位數 Median", "Q1", "Q3", "最小值", "最大值"]


# ============================================================ 資料結構
@dataclass
class Figure:
    caption: str
    png: bytes


@dataclass
class Section:
    title: str
    paragraphs: list[str] = field(default_factory=list)  # 可含 **粗體**
    tables: list = field(default_factory=list)  # [(名稱, DataFrame)]
    figures: list[Figure] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


@dataclass
class Report:
    title: str
    dataset: str
    created: str
    highlights: list[str] = field(default_factory=list)
    sections: list[Section] = field(default_factory=list)
    methods: list[str] = field(default_factory=list)
    software: str = ""
    lang: str = "zh"  # 產生報告時的語言；匯出時的標籤 (表、圖、重點摘要…) 跟著它

    def all_tables(self) -> list:
        return [(f"{i + 1}.{j + 1} {name}", tab) for i, s in enumerate(self.sections) for j, (name, tab) in enumerate(s.tables)]


@dataclass
class ReportOptions:
    outcomes: list[str]
    group: str | None = None
    categorical: list[str] = field(default_factory=list)
    strategy: str = "auto"
    correlation: bool = True
    figures: bool = True
    title: str = ""
    dpi: int = 150


# ============================================================ 小工具
def p_str(p: float) -> str:
    if p is None or not np.isfinite(p):
        return "p = n/a"
    return "p < 0.001" if p < 0.001 else f"p = {p:.3f}"


def _sig(p: float) -> bool:
    return bool(p is not None and np.isfinite(p) and p < ALPHA)


def _num(v: float) -> str:
    if v is None or not np.isfinite(v):
        return "n/a"
    a = abs(v)
    if a != 0 and (a < 0.001 or a >= 1e6):
        return f"{v:.3g}"
    return f"{v:.1f}" if a >= 100 else f"{v:.2f}" if a >= 1 else f"{v:.3f}"


def _df_txt(v: float) -> str:
    return str(int(round(v))) if abs(v - round(v)) < 1e-9 else f"{v:.1f}"


def _join(items) -> str:
    return _("、").join(str(x) for x in items)


def _style() -> Style:
    return Style(width=5.2, height=3.9, title_size=12, label_size=11, tick_size=9.5, legend_size=9)


def _png(fig, rc: dict, dpi: int) -> bytes:
    try:
        content, _ext, _mime = export.fig_bytes(fig, "PNG", dpi=dpi, rc=rc)
    finally:
        export.close(fig)
    return content


def _chart(key: str, df: pd.DataFrame, m: dict, o: dict, style: Style, dpi: int, **kw) -> bytes:
    res = plots.render(key, df, m, o, style, **kw)
    return _png(res.fig, res.rc, dpi)


def _grid(n: int, ncols: int = 3, w: float = 3.7, h: float = 3.0):
    ncols = max(1, min(ncols, n))
    nrows = math.ceil(n / ncols)
    if n == 1:  # 單張圖加寬，右側圖例才不會把圖擠窄
        w, h = 5.4, 3.6
    fig, axs = plt.subplots(nrows, ncols, figsize=(w * ncols, h * nrows), squeeze=False, constrained_layout=True)
    for ax in axs.flat[n:]:
        ax.remove()
    return fig, list(axs.flat[:n])


def _rc(style: Style) -> dict:
    return build_rc(style, fonts.ensure_cjk_font() if style.use_cjk else None)


def _legend_right(fig, axes: list) -> None:
    """多面板圖：移除各子圖的圖例，改成整張圖右側一個共用圖例，不壓到資料。"""
    src = next((ax.get_legend() for ax in axes if ax.get_legend() is not None), None)
    if src is None:
        return
    handles = getattr(src, "legend_handles", None) or src.legendHandles
    labels = [t.get_text() for t in src.get_texts()]
    title = src.get_title().get_text()
    for ax in axes:
        if ax.get_legend() is not None:
            ax.get_legend().remove()
    fig.legend(handles, labels, title=title or None, loc="outside right center", frameon=False)


def _shapiro(x: np.ndarray) -> float:
    if len(x) < 3 or np.ptp(x) == 0:
        return np.nan
    if len(x) > 5000:  # Shapiro-Wilk 上限；大樣本改用 D'Agostino K²
        return float(sps.normaltest(x).pvalue)
    return float(sps.shapiro(x).pvalue)


def _choose(strategy: str, normal: bool | None, homog_p: float, k: int) -> str:
    parametric = strategy == "parametric" or (strategy == "auto" and normal is not False)
    if not parametric:
        return "mwu" if k == 2 else "kw"
    unequal = np.isfinite(homog_p) and homog_p < ALPHA
    if k == 2:
        return "welch" if unequal else "student"
    return "welch_anova" if unequal else "anova"


def _short(method: str) -> str:
    """方法簡稱 (去掉「+ 事後檢定」)。"""
    return _(METHOD_LABEL[method]).split(" + ")[0]


# ============================================================ 各章節
def _overview(df: pd.DataFrame, work: pd.DataFrame, opts: ReportOptions, levels: list, dataset: str) -> Section:
    sec = Section(_("資料概覽"))
    cols = list(dict.fromkeys(opts.outcomes + ([opts.group] if opts.group else []) + opts.categorical))
    sec.paragraphs.append(_("資料集「{name}」共有 {rows:,} 筆資料、{cols} 個欄位。本報告分析 {k} 個數值結果變項：{vars}。",
                            name=dataset, rows=len(df), cols=df.shape[1], k=len(opts.outcomes), vars=_join(opts.outcomes)))
    if opts.group:
        counts = work[opts.group].value_counts().reindex(levels)
        sec.paragraphs.append(_("分組變項為「{g}」，共 {k} 組：{groups}。", g=opts.group, k=len(levels),
                                groups=_join(f"{lv} (n = {int(c)})" for lv, c in counts.items())))
        dropped = int(df[opts.group].isna().sum())
        if dropped:
            sec.warnings.append(_("有 {n} 筆資料的分組變項為空白，已排除於組間比較之外。", n=dropped))
        sec.tables.append((_("各組樣本數"), pd.DataFrame({"組別": levels, "n": counts.to_numpy(int),
                                                       "比例 (%)": counts.to_numpy(float) / counts.sum() * 100})))
    sec.tables.append((_("變項摘要"), data.summary_table(df[cols])))
    for v in opts.outcomes:
        miss = int(pd.to_numeric(df[v], errors="coerce").isna().sum())
        if miss:
            sec.warnings.append(_("「{v}」有 {n} 筆缺失或非數值資料，分析時已排除。", v=v, n=miss))
    return sec


def _descriptive(work: pd.DataFrame, opts: ReportOptions, levels: list, style: Style) -> Section:
    sec = Section(_("描述性統計"))
    res = bs.descriptive(work, opts.outcomes, opts.group)
    tab = res.tables[0][1]
    sec.tables.append((_("描述性統計"), tab[[c for c in _DESC_COLS if c in tab.columns]]))
    sec.paragraphs.append(_("連續變項以平均值 ± 標準差 (SD)、平均值的 95% 信賴區間，以及中位數與四分位數 (Q1–Q3) 描述。"))
    if opts.figures:
        rc = _rc(style)
        with plots.PLOT_LOCK, plt.rc_context(rc):
            fig, axes = _grid(len(opts.outcomes))
            pal = style.colors(max(len(levels), 1))
            for ax, v in zip(axes, opts.outcomes):
                sns.histplot(data=work, x=v, hue=opts.group, hue_order=levels or None, kde=True, stat="density",
                             common_norm=False, element="step", alpha=0.25, palette=pal if opts.group else None,
                             color=None if opts.group else pal[0], ax=ax)
                ax.set_title(v)
                ax.set_ylabel(_("密度"))
                sns.despine(ax=ax)
            _legend_right(fig, axes)
        cap = _("各結果變項的分布 (直方圖 + 核密度估計，依組別著色)") if opts.group else _("各結果變項的分布 (直方圖 + 核密度估計)")
        sec.figures.append(Figure(cap, _png(fig, rc, opts.dpi)))
    return sec


def _assumptions(work: pd.DataFrame, opts: ReportOptions, levels: list, style: Style) -> tuple[Section, dict]:
    """常態性與變異數同質性；回傳章節與 {變項: (是否常態, 同質性 p)}。"""
    sec = Section(_("前提檢定：常態性與變異數同質性"))
    rows, var_rows, verdict = [], [], {}
    for v in opts.outcomes:
        samples = bs.group_samples(work, v, opts.group, levels) if opts.group else {"全部": bs._clean(work[v])}
        ps = {}
        for lv, x in samples.items():
            p = _shapiro(x)
            ps[lv] = p
            rows.append({"變項": v, "組別": lv, "n": len(x), "Shapiro-Wilk p": p,
                         "結論": "—" if not np.isfinite(p) else ("符合常態" if p >= ALPHA else "偏離常態")})
        finite = [p for p in ps.values() if np.isfinite(p)]
        normal = None if not finite else all(p >= ALPHA for p in finite)
        homog_p = np.nan
        usable = [x for x in samples.values() if len(x) >= 2]
        if opts.group and len(usable) >= 2:
            lev = sps.levene(*usable, center="median")
            homog_p = float(lev.pvalue)
            var_rows.append({"變項": v, "Brown-Forsythe F": float(lev.statistic), "p 值": homog_p,
                             "結論": "變異數同質" if homog_p >= ALPHA else "變異數不同質"})
        verdict[v] = (normal, homog_p)
    sec.tables.append((_("常態性檢定 (Shapiro-Wilk)"), pd.DataFrame(rows)))
    if var_rows:
        sec.tables.append((_("變異數同質性檢定 (Brown-Forsythe)"), pd.DataFrame(var_rows)))
    off = [v for v, (n, _h) in verdict.items() if n is False]
    if off:
        sec.paragraphs.append(_("Shapiro-Wilk 檢定顯示 {vars} 至少有一組顯著偏離常態分布 (p < {a})。", vars=_join(off), a=ALPHA))
    else:
        sec.paragraphs.append(_("各結果變項在各組皆未顯著偏離常態分布。"))
    hetero = [r["變項"] for r in var_rows if r["p 值"] < ALPHA]
    if var_rows:
        sec.paragraphs.append(_("Brown-Forsythe 檢定顯示 {vars} 的組間變異數不同質。", vars=_join(hetero)) if hetero
                              else _("各結果變項的組間變異數皆可視為同質。"))
    sec.paragraphs.append(_("注意：樣本數小時常態檢定的檢定力低，樣本數大時則過於敏感，建議同時參考下方 Q-Q 圖。"))
    if any(len(x) < 3 for v in opts.outcomes for x in
           (bs.group_samples(work, v, opts.group, levels) if opts.group else {"_": bs._clean(work[v])}).values()):
        sec.warnings.append(_("部分組別樣本數 < 3，無法檢定常態性，結果請保守解讀。"))
    if opts.figures:
        rc = _rc(style)
        with plots.PLOT_LOCK, plt.rc_context(rc):
            fig, axes = _grid(len(opts.outcomes))
            pal = style.colors(max(len(levels), 1))
            for ax, v in zip(axes, opts.outcomes):
                samples = bs.group_samples(work, v, opts.group, levels) if opts.group else {"全部": bs._clean(work[v])}
                for i, (lv, x) in enumerate(samples.items()):
                    if len(x) < 3:
                        continue
                    (osm, osr), (slope, icpt, _r) = sps.probplot(x, dist="norm")
                    sns.scatterplot(x=osm, y=osr, color=pal[i % len(pal)], s=14, alpha=0.8, edgecolor="none",
                                    label=str(lv) if opts.group else None, ax=ax)
                    ax.plot(osm, slope * osm + icpt, color=pal[i % len(pal)], lw=1)
                ax.set(title=v, xlabel=_("理論分位數"), ylabel=_("樣本分位數"))
                if opts.group and ax.get_legend_handles_labels()[0]:
                    ax.legend(frameon=False, fontsize=8, title=opts.group, title_fontsize=8)
                sns.despine(ax=ax)
            _legend_right(fig, axes)
        sec.figures.append(Figure(_("常態 Q-Q 圖 (點越貼近直線，越符合常態分布)"), _png(fig, rc, opts.dpi)))
    return sec, verdict


def _cell(x: np.ndarray, nonparam: bool) -> str:
    if len(x) == 0:
        return "—"
    if nonparam:
        q1, med, q3 = np.percentile(x, [25, 50, 75])
        return f"{_num(med)} [{_num(q1)}, {_num(q3)}]"
    return f"{_num(np.mean(x))} ± {_num(np.std(x, ddof=1)) if len(x) > 1 else 'n/a'}"


def _compare_one(work: pd.DataFrame, v: str, g: str, levels: list, verdict: tuple, strategy: str) -> dict | None:
    samples = {k: x for k, x in bs.group_samples(work, v, g, levels).items() if len(x) >= 2}
    k = len(samples)
    if k < 2:
        return None
    normal, homog_p = verdict
    method = _choose(strategy, normal, homog_p, k)
    keys = list(samples)
    out = {"method": method, "samples": samples, "posthoc": None, "sig_pairs": []}
    if k == 2:
        r = bs.two_group(work, v, g, keys[0], keys[1])
        t = dict(r.tables)
        row = t["檢定結果"].set_index("檢定").loc[_TEST_ROW[method]]
        eff = t["效果量"].set_index("效果量")["估計值"]
        stat, p = float(row["統計量"]), float(row["p 值"])
        if method == "mwu":
            out.update(stat_txt=f"U = {stat:.1f}", es_name="rank-biserial r", es=float(eff["Rank-biserial r"]),
                       diff_txt=_("中位數差 ({a} − {b}) = {d}。", a=keys[0], b=keys[1], d=_num(row["差異 (平均/中位數)"])))
        else:
            out.update(stat_txt=f"t({_df_txt(row['自由度'])}) = {stat:.3f}", es_name="Cohen's d", es=float(eff["Cohen's d"]),
                       diff_txt=_("平均差 ({a} − {b}) = {d}，95% CI [{lo}, {hi}]。", a=keys[0], b=keys[1],
                                  d=_num(row["差異 (平均/中位數)"]), lo=_num(row["CI 下限"]), hi=_num(row["CI 上限"])))
        out["p"] = p
    else:
        r = bs.one_way(work, v, g, posthoc=_POSTHOC[method], order=keys)
        t = dict(r.tables)
        row = t["整體檢定"].set_index("檢定").loc[_TEST_ROW[method]]
        stat, p = float(row["統計量"]), float(row["p 值"])
        n_tot = sum(len(x) for x in samples.values())
        if method == "kw":
            out.update(stat_txt=f"H({_df_txt(row['df1'])}) = {stat:.3f}", es_name="ε²", es=stat * (n_tot + 1) / (n_tot ** 2 - 1))
        else:
            eta2 = float(t["ANOVA 表"]["η²"].iloc[0])
            out.update(stat_txt=f"F({_df_txt(row['df1'])}, {_df_txt(row['df2'])}) = {stat:.3f}", es_name="η²", es=eta2)
        ph = t[f"事後檢定：{_POSTHOC[method]}"]
        out["posthoc"] = ph
        out["sig_pairs"] = [(a, b, q) for a, b, q in zip(ph["group1"], ph["group2"], ph["p_adj"]) if _sig(q)]
        out["p"] = p
        out["diff_txt"] = ""
    return out


_WHY = {"user": "依設定一律採用{method}。", "nonnormal": "因有組別偏離常態，採用{method}。",
        "unequal": "因變異數不同質，採用{method}。", "ok": "因資料符合常態且變異數同質，採用{method}。"}


def _comparisons(work: pd.DataFrame, opts: ReportOptions, levels: list, verdict: dict, style: Style,
                 highlights: list, used: set) -> Section:
    g = opts.group
    sec = Section(_("組間比較：依「{g}」分組", g=g))
    summary, table1 = [], []
    for v in opts.outcomes:
        r = _compare_one(work, v, g, levels, verdict[v], opts.strategy)
        if r is None:
            sec.warnings.append(_("「{v}」有效組別不足兩組 (每組至少需 2 筆)，略過比較。", v=v))
            continue
        method, p = r["method"], r["p"]
        used.add(method)
        nonparam = method in ("mwu", "kw")
        normal, homog_p = verdict[v]
        why = ("user" if opts.strategy != "auto" else "nonnormal" if normal is False else
               "unequal" if np.isfinite(homog_p) and homog_p < ALPHA else "ok")
        sig = _sig(p)
        stat = _("{stat}，{p}，{es_name} = {es}", stat=r["stat_txt"], p=p_str(p), es_name=r["es_name"], es=_num(r["es"]))
        text = (_("**{v}**：", v=v) + _(_WHY[why], method=_(METHOD_LABEL[method])) +
                (_("結果顯示各組間有顯著差異 ({stat})。", stat=stat) if sig else _("結果顯示各組間無顯著差異 ({stat})。", stat=stat)))
        text += r.get("diff_txt") or ""
        if r["posthoc"] is not None and sig:
            text += (_("事後檢定達顯著的組合：{pairs}。", pairs=_join(f"{a} vs {b} ({p_str(q)})" for a, b, q in r["sig_pairs"]))
                     if r["sig_pairs"] else _("事後檢定未發現任何兩組間有顯著差異。"))
        sec.paragraphs.append(text)
        hl = (_("{v}：各組間有顯著差異 ({method}，{p})", v=v, method=_short(method), p=p_str(p)) if sig else
              _("{v}：各組間無顯著差異 ({method}，{p})", v=v, method=_short(method), p=p_str(p)))
        if r["sig_pairs"]:
            hl += _("；顯著組合 {pairs}", pairs=_join(f"{a} vs {b}" for a, b, _q in r["sig_pairs"]))
        highlights.append(hl)
        summary.append({"變項": v, "組數": len(r["samples"]), "總 n": sum(len(x) for x in r["samples"].values()),
                        "檢定方法": _short(method), "統計量": r["stat_txt"], "p 值": p,
                        "效果量": f"{r['es_name']} = {_num(r['es'])}", "結論": "顯著" if sig else "不顯著"})
        cells = {lv: _cell(r["samples"].get(lv, np.array([])), nonparam) for lv in levels}
        table1.append({"變項": v, "呈現方式": "中位數 [Q1, Q3]" if nonparam else "平均值 ± SD", **cells,
                       "檢定": _short(method), "p 值": p})
        if r["posthoc"] is not None and sig:  # 整體檢定不顯著時不解讀事後檢定
            sec.tables.append((_("{v} 事後檢定 ({ph})", v=v, ph=_POSTHOC[method]), r["posthoc"]))
        if opts.figures:
            # 三組以上且整體檢定不顯著時不畫事後比較括號，避免與文字結論矛盾
            brackets = sig or len(r["samples"]) == 2
            sig_opt = dict(enabled=brackets, test=_SIG_TEST[method], pairs_mode="all", ref=None,
                           correction="holm" if method == "kw" else None, display="星號 (*)",
                           hide_ns=len(r["samples"]) > 3, font_size=None, line_width=1.0)
            try:
                png = _chart("box", work, {"x": g, "y": v}, {"points": "散點 (strip)"}, style, opts.dpi,
                             orders={g: [lv for lv in levels if lv in r["samples"]]}, sig=sig_opt)
                cap = _("{v} 依 {g} 分組的盒鬚圖 + 個別資料點", v=v, g=g)
                if brackets:
                    note = _("，只標示顯著的組合") if sig_opt["hide_ns"] else ""
                    cap += _("；括號為 {ph} 結果{note} (ns p ≥ 0.05、* p < 0.05、** p < 0.01、*** p < 0.001、**** p < 0.0001)",
                             ph=_POSTHOC.get(method) or _(METHOD_LABEL[method]), note=note)
                else:
                    cap += _("；整體檢定未達顯著，不進行事後比較")
                sec.figures.append(Figure(cap, png))
            except Exception as err:
                sec.warnings.append(_("「{v}」圖表繪製失敗：{err}", v=v, err=err))
    if summary:
        sec.tables.insert(0, (_("各組摘要與檢定結果 (Table 1)"), pd.DataFrame(table1)))
        sec.tables.insert(1, (_("組間比較結果"), pd.DataFrame(summary)))
    return sec


def _correlation(work: pd.DataFrame, opts: ReportOptions, style: Style, highlights: list, used: set) -> Section:
    vals = opts.outcomes
    pvals = [_shapiro(bs._clean(work[v])) for v in vals]
    normal = all(not np.isfinite(p) or p >= ALPHA for p in pvals)
    method = {"parametric": "pearson", "nonparametric": "spearman"}.get(opts.strategy, "pearson" if normal else "spearman")
    used.add(method)
    name = "Pearson" if method == "pearson" else "Spearman"
    sec = Section(_("相關分析 ({name})", name=name))
    res = bs.correlation(work, vals, method)
    pairs = res.tables[0][1]
    if len(pairs):
        # 欄位順序固定：0 變項1、1 變項2、2 n、3 r、4 CI 下限、5 CI 上限、6 p、7 p (FDR)、8 強度
        pairs = pairs[["變項 1", "變項 2", "n", "r", "CI 下限", "CI 上限", "p 值", "p (FDR)", "強度"]]
        sec.tables.append((_("成對相關"), pairs))
    why = ("各變項皆符合常態" if method == "pearson" else "有變項偏離常態") if opts.strategy == "auto" else "依設定"
    sec.paragraphs.append(_("{why}，以 {name} 相關係數評估結果變項兩兩之間的{kind}關聯，並以 Benjamini-Hochberg 法控制偽發現率 (FDR)。",
                            why=_(why), name=name, kind=_("線性" if method == "pearson" else "單調")))
    sig_rows = pairs[pairs["p (FDR)"] < ALPHA] if len(pairs) else pairs
    if len(sig_rows):
        for r in sig_rows.itertuples(index=False):
            sec.paragraphs.append(_("{a} 與 {b} 呈{dir}相關 (r = {r:.3f}，95% CI [{lo:.3f}, {hi:.3f}]，FDR 校正 {p}，n = {n})，屬{strength}相關。",
                                    a=r[0], b=r[1], dir=_("正" if r[3] > 0 else "負"), r=r[3], lo=r[4], hi=r[5],
                                    p=p_str(r[7]), n=r[2], strength=_(r[8])))
        highlights.append(_("顯著相關：{pairs}", pairs=_join(f"{r[0]}–{r[1]} (r = {r[3]:.2f})" for r in sig_rows.itertuples(index=False))))
    elif len(pairs):
        sec.paragraphs.append(_("經 FDR 校正後，沒有任何一對變項達顯著相關。"))
    if opts.figures:
        try:
            png = _chart("corr", work, {"values": vals}, {"method": method}, Style(width=4.8, height=4.2, title_size=12,
                                                                                 label_size=10, tick_size=9), opts.dpi)
            sec.figures.append(Figure(_("{name} 相關係數熱圖 (星號為未校正 p 值的顯著性)", name=name), png))
            if len(vals) <= 5:
                m = {"values": vals, "hue": opts.group} if opts.group else {"values": vals}
                png = _chart("pairplot", work, m, {"panel": 2.0, "corner": True}, style, opts.dpi,
                             orders={opts.group: plots.category_levels(work[opts.group])} if opts.group else None)
                sec.figures.append(Figure(_("配對散佈圖 (對角線為各變項分布)"), png))
        except Exception as err:
            sec.warnings.append(_("相關圖繪製失敗：{err}", err=err))
    return sec


def _categorical(work: pd.DataFrame, opts: ReportOptions, levels: list, style: Style, highlights: list, used: set) -> Section:
    g = opts.group
    cats = [c for c in opts.categorical if c != g]
    sec = Section(_("類別變項分析") + (_("：與「{g}」的關聯", g=g) if g else ""))
    if not g:
        for c in cats:
            vc = work[c].value_counts(dropna=False)
            sec.tables.append((_("{c} 次數分配", c=c), pd.DataFrame({"類別": vc.index.astype(str), "次數": vc.to_numpy(),
                                                                "百分比 (%)": vc.to_numpy() / vc.sum() * 100})))
        sec.paragraphs.append(_("未指定分組變項，僅列出各類別變項的次數分配。"))
        return sec
    summary = []
    for c in cats:
        sub = work.dropna(subset=[g, c]).copy()
        sub[c] = data.as_category_labels(sub[c])
        res = bs.contingency(sub, g, c)
        if res.warnings and not res.tables:
            sec.warnings.extend(_("「{c}」：{w}", c=c, w=w) for w in res.warnings)
            continue
        t = dict(res.tables)
        tests = t["檢定結果"].set_index("檢定")
        exp = t["期望次數"].set_index(g).to_numpy(float)
        two_by_two = exp.shape == (2, 2)
        use_fisher = two_by_two and (exp < 5).any()
        if use_fisher:
            p, test, stat_txt = float(tests.loc["Fisher 精確檢定", "p 值"]), "Fisher 精確檢定", ""
            used.add("fisher")
        else:
            row = tests.iloc[0]
            p, test = float(row["p 值"]), str(row.name)
            stat_txt = _("χ²({df}) = {x:.3f}，", df=int(row["df"]), x=row["統計量"])
            used.add("chi2")
        v = float(tests.loc["Cramér's V (效果量)", "統計量"])
        sig = _sig(p)
        kw = dict(c=c, g=g, test=_(test), stat=stat_txt, p=p_str(p), v=f"{v:.3f}")
        sec.paragraphs.append(_("**{c}**：以 {test} 評估「{g}」與「{c}」的關聯，結果達統計顯著 ({stat}{p}，Cramér's V = {v})。", **kw) if sig
                              else _("**{c}**：以 {test} 評估「{g}」與「{c}」的關聯，結果未達統計顯著 ({stat}{p}，Cramér's V = {v})。", **kw))
        short = _(test).split(" (")[0]
        highlights.append(_("{c} 與 {g}：有顯著關聯 ({test}，{p})", c=c, g=g, test=short, p=p_str(p)) if sig
                          else _("{c} 與 {g}：無顯著關聯 ({test}，{p})", c=c, g=g, test=short, p=p_str(p)))
        summary.append({"變項": c, "檢定": test, "p 值": p, "Cramér's V": v, "結論": "顯著" if sig else "不顯著"})
        sec.tables.append((_("{g} × {c} 次數", g=g, c=c), t["觀察次數"]))
        sec.tables.append((_("{g} × {c} 列百分比 (%)", g=g, c=c), t["列百分比 (%)"]))
        if "2×2 效果量" in t:
            sec.tables.append((_("{g} × {c} 勝算比 / 相對風險", g=g, c=c), t["2×2 效果量"]))
        sec.warnings.extend(_("「{c}」：{w}", c=c, w=w) for w in res.warnings)
        if opts.figures:
            pct = _("百分比 (%)")
            ct = pd.crosstab(sub[g], sub[c], normalize="index") * 100
            long = ct.reindex([lv for lv in levels if lv in ct.index]).reset_index().melt(id_vars=g, var_name=c, value_name=pct)
            rc = _rc(style)
            with plots.PLOT_LOCK, plt.rc_context(rc):
                fig, ax = plt.subplots(figsize=(style.width, style.height), constrained_layout=True)
                sns.barplot(data=long, x=g, y=pct, hue=c, palette=style.colors(long[c].nunique()), ax=ax)
                ax.set_title(_("{c} 在各 {g} 的組成", c=c, g=g))
                sns.move_legend(ax, "upper left", bbox_to_anchor=(1.01, 1), frameon=False, title=c)
                tidy_xticklabels(ax, style)
                sns.despine(ax=ax)
            sec.figures.append(Figure(_("{c} 在各 {g} 組內的百分比", c=c, g=g), _png(fig, rc, opts.dpi)))
    if summary:
        sec.tables.insert(0, (_("類別變項關聯摘要"), pd.DataFrame(summary)))
    return sec


def _methods(used: set, opts: ReportOptions) -> list[str]:
    out = [_("連續變項以平均值 ± 標準差 (SD) 或中位數 [第一四分位數, 第三四分位數] 描述；類別變項以次數與百分比描述。")]
    if opts.group:
        out.append(_("各組常態性以 Shapiro-Wilk 檢定評估，變異數同質性以 Brown-Forsythe 檢定 (以中位數為中心的 Levene 檢定) 評估。"))
        two = [_(METHOD_LABEL[m]) for m in ("student", "welch", "mwu") if m in used]
        many = [_(METHOD_LABEL[m]) for m in ("anova", "welch_anova", "kw") if m in used]
        rule = (_("符合常態且變異數同質者使用母數檢定、變異數不同質者使用 Welch 校正、偏離常態者使用無母數檢定")
                if opts.strategy == "auto" else
                _("依使用者設定一律使用母數檢定") if opts.strategy == "parametric" else _("依使用者設定一律使用無母數檢定"))
        if two:
            out.append(_("兩組比較使用 {tests} ({rule})；效果量以 Cohen's d 或 rank-biserial 相關係數表示。", tests=_join(two), rule=rule))
        if many:
            out.append(_("三組以上比較使用 {tests} ({rule})；效果量以 η² 或 ε² 表示，事後檢定已校正多重比較。", tests=_join(many), rule=rule))
    if used & {"pearson", "spearman"}:
        out.append(_("變項間關聯以 {name} 相關係數評估，95% 信賴區間以 Fisher z 轉換計算，並以 Benjamini-Hochberg 法校正多重比較。",
                     name="Pearson" if "pearson" in used else "Spearman"))
    if used & {"chi2", "fisher"}:
        out.append(_("類別變項的關聯以 Pearson 卡方檢定評估 (2×2 表使用 Yates 連續性校正)，"
                     "2×2 表中任一期望次數 < 5 時改用 Fisher 精確檢定；效果量以 Cramér's V 表示。"))
    out.append(_("所有檢定皆為雙尾，顯著水準 α = {a}。", a=ALPHA))
    return out


def software_line() -> str:
    libs = [f"Python {platform.python_version()}", f"pandas {pd.__version__}", f"NumPy {np.__version__}",
            f"SciPy {scipy.__version__}", f"statsmodels {statsmodels.__version__}", f"seaborn {sns.__version__}",
            f"matplotlib {matplotlib.__version__}"]
    return f"Super Xtat v{__version__}" + _("；") + _join(libs)


# ============================================================ 主函式
def build_report(df: pd.DataFrame, opts: ReportOptions, dataset: str = "") -> Report:
    """依目前語言 (i18n) 產生報告。"""
    outcomes = [v for v in dict.fromkeys(opts.outcomes) if v in df.columns and v != opts.group]
    if not outcomes:
        raise ValueError(_("請至少選擇一個數值結果變項"))
    opts.outcomes = outcomes
    dataset = dataset or _("資料")
    work = df.copy()
    for v in outcomes:
        work[v] = pd.to_numeric(work[v], errors="coerce")
    levels: list = []
    if opts.group:
        work[opts.group] = data.as_category_labels(work[opts.group])
        levels = plots.category_levels(work[opts.group])
        work = work.dropna(subset=[opts.group])
        if len(levels) < 2:
            raise ValueError(_("分組變項「{g}」至少需要 2 個類別", g=opts.group))
        if len(levels) > 20:
            raise ValueError(_("分組變項「{g}」有 {n} 個類別，太多了 (上限 20)，請改選其他欄位", g=opts.group, n=len(levels)))
    style = _style()
    rep = Report(opts.title or _("統計分析報告"), dataset, dt.datetime.now().strftime("%Y-%m-%d %H:%M"), lang=i18n.get_lang())
    used: set = set()
    rep.sections.append(_overview(df, work, opts, levels, dataset))
    rep.sections.append(_descriptive(work, opts, levels, style))
    assum, verdict = _assumptions(work, opts, levels, style)
    rep.sections.append(assum)
    if opts.group:
        rep.sections.append(_comparisons(work, opts, levels, verdict, style, rep.highlights, used))
    if opts.correlation and len(outcomes) >= 2:
        rep.sections.append(_correlation(work, opts, style, rep.highlights, used))
    if [c for c in opts.categorical if c != opts.group]:
        rep.sections.append(_categorical(work, opts, levels, style, rep.highlights, used))
    rep.methods = _methods(used, opts)
    rep.software = software_line()
    if not rep.highlights:
        rep.highlights.append(_("本報告僅包含描述性統計與前提檢定；選擇分組變項即可進行組間比較。"))
    for sec in rep.sections:  # 表格欄名與文字儲存格轉成目前語言 (中文模式不變)
        sec.tables = [(name, i18n.table(tab)) for name, tab in sec.tables]
    return rep


# ============================================================ 匯出
_BOLD = re.compile(r"\*\*(.+?)\*\*")

REPORT_CSS = """
:root{--ink:#111;--muted:#5f6368;--line:#e3e3e3;--accent:#3e6ae1;--soft:#f6f7f9}
*{box-sizing:border-box}
body{margin:0;background:#fff;color:var(--ink);font:15px/1.7 "Inter","Segoe UI","Microsoft JhengHei","PingFang TC","Noto Sans TC",sans-serif}
main{max-width:980px;margin:0 auto;padding:48px 28px 64px}
header{border-bottom:1px solid var(--ink);padding-bottom:22px;margin-bottom:28px}
.brand{font:600 12px/1 "Barlow","Inter",sans-serif;letter-spacing:.42em;text-transform:uppercase;color:var(--muted)}
h1{font:600 34px/1.2 "Barlow","Inter","Microsoft JhengHei",sans-serif;letter-spacing:.02em;margin:14px 0 8px}
.meta{color:var(--muted);font-size:13px}
h2{font:600 20px/1.3 "Barlow","Inter","Microsoft JhengHei",sans-serif;letter-spacing:.04em;margin:44px 0 12px;
   padding-top:18px;border-top:1px solid var(--line)}
h2 .no{color:var(--accent);margin-right:10px}
h3{font-size:14px;font-weight:600;color:var(--muted);letter-spacing:.06em;margin:26px 0 8px}
.highlights{background:var(--soft);border-left:3px solid var(--accent);padding:14px 22px;margin:8px 0 6px}
.highlights h2{border:0;margin:0 0 6px;padding:0;font-size:15px;letter-spacing:.2em;text-transform:uppercase}
.highlights ul{margin:0;padding-left:20px}
.warn{background:#fff6e5;border-left:3px solid #e8a100;padding:8px 14px;margin:10px 0;font-size:13.5px}
.tbl-wrap{overflow-x:auto;margin:6px 0 10px}
table{border-collapse:collapse;font-size:12.5px;width:auto;min-width:50%}
th{font-weight:600;text-align:left;border-bottom:1.5px solid var(--ink);border-top:1.5px solid var(--ink);padding:6px 10px;white-space:nowrap}
td{padding:5px 10px;border-bottom:1px solid var(--line);white-space:nowrap;font-variant-numeric:tabular-nums}
tr:last-child td{border-bottom:1.5px solid var(--ink)}
figure{margin:18px 0 26px;text-align:center}
figure img{max-width:100%;height:auto}
figcaption{font-size:12.5px;color:var(--muted);margin-top:6px;text-align:left}
footer{margin-top:48px;padding-top:14px;border-top:1px solid var(--line);color:var(--muted);font-size:12px}
@media print{main{padding:0}h2{break-after:avoid}figure,table{break-inside:avoid}}
"""


def _inline(text: str) -> str:
    return _BOLD.sub(r"<strong>\1</strong>", html.escape(text))


def _labeler(rep: Report):
    """匯出用的標籤翻譯：依報告產生時的語言，而不是目前介面語言。"""
    return lambda text, **kw: i18n.tr(rep.lang, text, **kw)


def to_html(rep: Report) -> str:
    L = _labeler(rep)
    lang = "en" if rep.lang == "en" else "zh-Hant"
    parts = [f'<!doctype html><html lang="{lang}"><head><meta charset="utf-8">'
             f'<meta name="viewport" content="width=device-width,initial-scale=1"><title>{html.escape(rep.title)}</title>'
             '<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Barlow:wght@500;600'
             '&family=Inter:wght@400;600&display=swap">'
             f'<style>{REPORT_CSS}</style></head><body><main>',
             f'<header><div class="brand">Super Xtat</div><h1>{html.escape(rep.title)}</h1>'
             '<div class="meta">' + html.escape(L("資料：{d} · 產生時間：{c}", d=rep.dataset, c=rep.created)) + '</div></header>',
             f'<section class="highlights"><h2>{L("重點摘要")}</h2><ul>' +
             "".join(f"<li>{_inline(h)}</li>" for h in rep.highlights) + "</ul></section>"]
    for i, sec in enumerate(rep.sections, 1):
        parts.append(f'<section><h2><span class="no">{i:02d}</span>{html.escape(sec.title)}</h2>')
        parts += [f"<p>{_inline(p)}</p>" for p in sec.paragraphs]
        parts += [f'<div class="warn">{_inline(w)}</div>' for w in sec.warnings]
        for j, (name, tab) in enumerate(sec.tables, 1):
            body = data.format_table(tab).to_html(index=False, border=0, na_rep="", escape=True)
            parts.append(f'<h3>{L("表 {no}", no=f"{i}.{j}")}　{html.escape(name)}</h3><div class="tbl-wrap">{body}</div>')
        for j, f in enumerate(sec.figures, 1):
            b64 = base64.b64encode(f.png).decode()
            parts.append(f'<figure><img src="data:image/png;base64,{b64}" alt="{html.escape(f.caption)}">'
                         f'<figcaption>{L("圖 {no}", no=f"{i}.{j}")}　{html.escape(f.caption)}</figcaption></figure>')
        parts.append("</section>")
    parts.append(f'<section><h2><span class="no">{len(rep.sections) + 1:02d}</span>{L("統計方法")}</h2>' +
                 "".join(f"<p>{_inline(m)}</p>" for m in rep.methods) + "</section>")
    parts.append(f"<footer>{html.escape(rep.software)}</footer></main></body></html>")
    return "".join(parts)


def _cell_text(v) -> str:
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return ""
    return str(v)


def to_docx(rep: Report) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    for sname in ("Normal", "Heading 1", "Heading 2", "Heading 3", "Title", "Caption"):
        s = doc.styles[sname]
        s.font.name = "Calibri"
        s.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft JhengHei")
    doc.styles["Normal"].font.size = Pt(10.5)
    sect = doc.sections[0]
    sect.left_margin = sect.right_margin = Inches(0.9)

    def para(text: str, style: str | None = None):
        p = doc.add_paragraph(style=style)
        for k, chunk in enumerate(re.split(r"\*\*", text)):
            if chunk:
                p.add_run(chunk).bold = bool(k % 2)
        return p

    brand = doc.add_paragraph()
    run = brand.add_run("S U P E R   X T A T")
    run.font.size, run.font.color.rgb = Pt(9), RGBColor(0x5F, 0x63, 0x68)
    L = _labeler(rep)
    doc.add_heading(rep.title, level=0)
    para(L("資料：{d} · 產生時間：{c}", d=rep.dataset, c=rep.created))
    doc.add_heading(L("重點摘要"), level=1)
    for h in rep.highlights:
        para(h, "List Bullet")
    width = sect.page_width - sect.left_margin - sect.right_margin
    for i, sec in enumerate(rep.sections, 1):
        doc.add_heading(f"{i}. {sec.title}", level=1)
        for p in sec.paragraphs:
            para(p)
        for w in sec.warnings:
            para(f"⚠ {w}").runs[0].font.color.rgb = RGBColor(0xB0, 0x6A, 0x00)
        for j, (name, tab) in enumerate(sec.tables, 1):
            doc.add_heading(L("表 {no}", no=f"{i}.{j}") + f"　{name}", level=3)
            shown = data.format_table(tab)
            t = doc.add_table(rows=1, cols=len(shown.columns))
            t.style = "Light List"
            for c, col in enumerate(shown.columns):
                t.rows[0].cells[c].text = str(col)
            for _i, r in shown.iterrows():
                cells = t.add_row().cells
                for c, v in enumerate(r):
                    cells[c].text = _cell_text(v)
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for rn in p.runs:
                            rn.font.size = Pt(8.5)
        for j, f in enumerate(sec.figures, 1):
            doc.add_picture(io.BytesIO(f.png), width=min(width, Inches(6.2)))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap = doc.add_paragraph(L("圖 {no}", no=f"{i}.{j}") + f"　{f.caption}", style="Caption")
            cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_heading(f"{len(rep.sections) + 1}. " + L("統計方法"), level=1)
    for m in rep.methods:
        para(m)
    foot = doc.add_paragraph(rep.software)
    foot.runs[0].font.size, foot.runs[0].font.color.rgb = Pt(8), RGBColor(0x5F, 0x63, 0x68)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------ PowerPoint
# 簡報風格：顏色 + 版面開關。kicker = 頁首英文小標、bar = 標題下強調線、dividers = 每章前加章節分隔頁
# appbar = 頂部色帶標題列 (Material)、rules = 三線表 (學術)、cover = left / center / band / academic
_THEME_BASE = dict(on_accent="FFFFFF", kicker=False, bar=False, dividers=False, appbar=False, rules=False,
                   title_spacing=0, title_y=0.5, title_color="ink", body_size=18)
PPT_THEMES: dict[str, dict] = {
    "Super Xtat 科技風": {**_THEME_BASE, **dict(
        bg="000000", panel="0B0C0E", head="1A1C20", band="111214", ink="F2F3F5", dim="8A8F98", accent="3E6AE1",
        warn="E8A100", head_ink="dim", latin="Segoe UI", ea="Microsoft JhengHei",
        kicker=True, bar=True, title_size=26, title_bold=False, title_spacing=100, title_y=0.62, cover="left", body_size=17)},
    "Apple 簡約風": {**_THEME_BASE, **dict(
        bg="FFFFFF", panel="FFFFFF", head="F5F5F7", band="FBFBFD", ink="1D1D1F", dim="6E6E73", accent="0071E3",
        warn="B25000", head_ink="ink", latin="Segoe UI", ea="Microsoft JhengHei",
        dividers=True, title_size=30, title_bold=True, cover="center")},
    "學術白底": {**_THEME_BASE, **dict(
        bg="FFFFFF", panel="FFFFFF", head="FFFFFF", band="FFFFFF", ink="111111", dim="555555", accent="1F3A5F",
        warn="8A4B00", head_ink="ink", latin="Times New Roman", ea="PMingLiU",
        bar=True, rules=True, title_size=28, title_bold=True, title_color="accent", cover="academic", body_size=17)},
    "Google Material": {**_THEME_BASE, **dict(
        bg="FFFFFF", panel="FFFFFF", head="E8F0FE", band="F8F9FA", ink="202124", dim="5F6368", accent="1A73E8",
        warn="B06000", head_ink="accent", latin="Segoe UI", ea="Microsoft JhengHei",
        appbar=True, title_size=26, title_bold=False, title_y=0.3, title_color="on_accent", cover="band")},
}
DEFAULT_PPT_THEME = "Super Xtat 科技風"
_SLIDE_W, _SLIDE_H = 13.333, 7.5  # 16:9 (英吋)
_TABLE_ROWS = 12  # 每張投影片最多幾列資料
_TEXT_CHARS = 700  # 每張文字投影片大約容納的字數


def _ppt_font(run, t: dict, size: float, color: str = "ink", bold: bool = False, spacing: int | None = None) -> None:
    from pptx.dml.color import RGBColor
    from pptx.oxml.ns import qn
    from pptx.util import Pt

    f = run.font
    f.size, f.bold, f.name = Pt(size), bold, t["latin"]
    f.color.rgb = RGBColor.from_string(t.get(color, color))
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = rpr.makeelement(qn("a:ea"), {})
        rpr.append(ea)
    ea.set("typeface", t["ea"])
    if spacing:
        rpr.set("spc", str(spacing))


def _ppt_text(slide, t: dict, x, y, w, h, paras: list, size: float = 16, color: str = "ink", gap: float = 6,
              markup: bool = True, bold: bool = False, center: bool = False):
    """paras：字串 (markup=True 時可含 **粗體**) 或 (字串, 顏色)。"""
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(paras):
        text, col = item if isinstance(item, tuple) else (item, color)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        if center:
            p.alignment = PP_ALIGN.CENTER
        for k, chunk in enumerate(re.split(r"\*\*", text) if markup else [text]):
            if chunk:
                run = p.add_run()
                run.text = chunk
                _ppt_font(run, t, size, col, bold=bold or bool(k % 2))
    return box


def _ppt_blank(prs, t: dict):
    from pptx.dml.color import RGBColor
    from pptx.util import Inches

    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string(t["bg"])
    tf = s.shapes.add_textbox(Inches(_SLIDE_W - 1.6), Inches(_SLIDE_H - 0.5), Inches(1.0), Inches(0.3)).text_frame
    run = tf.paragraphs[0].add_run()
    run.text = f"{len(prs.slides):02d}"
    _ppt_font(run, t, 10, "dim", spacing=200)
    return s


def _ppt_bar(slide, t: dict, x: float, y: float, w: float = 1.3, h: float = 0.035, color: str = "accent") -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor.from_string(t.get(color, color))
    bar.line.fill.background()
    bar.shadow.inherit = False


def _ppt_slide(prs, t: dict, kicker: str, title: str):
    s = _ppt_blank(prs, t)
    if t["appbar"]:
        _ppt_bar(s, t, 0, 0, _SLIDE_W, 1.15)
    if t["kicker"]:
        _ppt_text(s, t, 0.6, 0.32, 12, 0.35, [f"SUPER XTAT  ·  {kicker}".upper()], size=10, color="dim", markup=False)
        s.shapes[-1].text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr().set("spc", "400")
    _ppt_text(s, t, 0.6, t["title_y"], 12.1, 0.8, [title], size=t["title_size"], color=t["title_color"], markup=False,
              bold=t["title_bold"])
    if t["title_spacing"]:
        s.shapes[-1].text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr().set("spc", str(t["title_spacing"]))
    if t["bar"]:
        _ppt_bar(s, t, 0.62, 1.38, w=_SLIDE_W - 1.24 if t["rules"] else 1.3, h=0.02 if t["rules"] else 0.035)
    return s


def _ppt_cover(prs, t: dict, rep: Report) -> None:
    s = _ppt_blank(prs, t)
    sub = f"{rep.dataset}　·　{rep.created}"
    if t["cover"] == "band":  # Material：上半部整塊主色
        _ppt_bar(s, t, 0, 0, _SLIDE_W, 4.6)
        _ppt_text(s, t, 0.8, 0.6, 8, 0.4, ["Super Xtat · Statistical Report"], size=14, color="on_accent", markup=False)
        _ppt_text(s, t, 0.8, 2.6, _SLIDE_W - 1.6, 1.6, [rep.title], size=48, color="on_accent", markup=False)
        _ppt_text(s, t, 0.8, 4.95, _SLIDE_W - 1.6, 0.8, [sub], size=18, color="dim", markup=False)
        return
    if t["cover"] == "academic":  # 學術：置中標題，上下細線
        _ppt_bar(s, t, 1.5, 2.35, _SLIDE_W - 3.0, 0.02, color="ink")
        _ppt_text(s, t, 1.0, 2.6, _SLIDE_W - 2.0, 1.4, [rep.title], size=40, bold=True, center=True, markup=False)
        _ppt_bar(s, t, 1.5, 4.05, _SLIDE_W - 3.0, 0.02, color="ink")
        _ppt_text(s, t, 1.0, 4.3, _SLIDE_W - 2.0, 0.8, [sub], size=16, color="dim", center=True, markup=False)
        return
    if t["cover"] == "center":
        _ppt_text(s, t, 0.8, 2.55, _SLIDE_W - 1.6, 1.4, [rep.title], size=54, bold=True, center=True, markup=False)
        _ppt_text(s, t, 0.8, 4.05, _SLIDE_W - 1.6, 0.9, [f"{rep.dataset}　·　{rep.created}"], size=18, color="dim",
                  center=True, markup=False)
        return
    _ppt_text(s, t, 0.6, 0.32, 12, 0.35, ["SUPER XTAT  ·  STATISTICAL REPORT"], size=10, color="dim", markup=False)
    s.shapes[-1].text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr().set("spc", "400")
    _ppt_text(s, t, 0.6, 2.6, 12, 1.4, [rep.title], size=44, markup=False)
    _ppt_bar(s, t, 0.62, 3.95)
    L = _labeler(rep)
    _ppt_text(s, t, 0.6, 4.2, 12, 1.0, [L("資料：{d}", d=rep.dataset), L("產生時間：{c}", c=rep.created)], size=16,
              color="dim", gap=4, markup=False)


def _ppt_divider(prs, t: dict, no: str, title: str) -> None:
    s = _ppt_blank(prs, t)
    _ppt_text(s, t, 0.8, 2.6, _SLIDE_W - 1.6, 0.6, [no], size=20, color="accent", bold=True, center=True, markup=False)
    _ppt_text(s, t, 0.8, 3.2, _SLIDE_W - 1.6, 1.4, [title], size=44, bold=True, center=True, markup=False)


def _chunks(paras: list, limit: int = _TEXT_CHARS) -> list:
    out, cur, n = [], [], 0
    for p in paras:
        size = len(p[0] if isinstance(p, tuple) else p)
        if cur and n + size > limit:
            out.append(cur)
            cur, n = [], 0
        cur.append(p)
        n += size
    if cur:
        out.append(cur)
    return out


def _ppt_table(slide, t: dict, tab: pd.DataFrame, top: float = 1.7) -> None:
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    shown = data.format_table(tab)
    rows, cols = len(shown) + 1, len(shown.columns)
    size = 12 if cols <= 6 else 10.5 if cols <= 10 else 9
    cells = [[str(c) for c in shown.columns]] + [[_cell_text(v) for v in r] for r in shown.itertuples(index=False)]
    width = _SLIDE_W - 1.2
    weights = [max(4, min(28, max(len(r[c]) for r in cells))) for c in range(cols)]
    shape = slide.shapes.add_table(rows, cols, Inches(0.6), Inches(top), Inches(width), Inches(0.36 * rows))
    tbl = shape.table
    tbl.first_row, tbl.horz_banding = True, False
    for c, wgt in enumerate(weights):
        tbl.columns[c].width = Inches(width * wgt / sum(weights))
    for r in range(rows):
        tbl.rows[r].height = Inches(0.34)
        for c in range(cols):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(t["head"] if r == 0 else t["band"] if r % 2 else t["panel"])
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Pt(2)
            tf = cell.text_frame
            tf.word_wrap = True
            run = tf.paragraphs[0].add_run()
            run.text = cells[r][c]
            _ppt_font(run, t, size, "ink" if r else t["head_ink"], bold=r == 0)
            if t["rules"]:  # 三線表：表頭上下各一線、表格底一線
                _cell_rules(cell, t, top=1.5 if r == 0 else 0, bottom=0.75 if r == 0 else 1.5 if r == rows - 1 else 0)


def _cell_rules(cell, t: dict, top: float = 0, bottom: float = 0) -> None:
    """加上儲存格上下框線 (pt)；0 = 不設定 (沿用表格樣式，白底上看不見)。a:lnT / a:lnB 依 schema 須放在填色之前。"""
    from pptx.oxml.ns import qn

    tcpr = cell._tc.get_or_add_tcPr()
    for tag, pt in (("a:lnB", bottom), ("a:lnT", top)):  # 依序插到最前面 → 最終順序為 lnT、lnB
        if not pt:
            continue
        ln = tcpr.makeelement(qn(tag), {"w": str(int(pt * 12700))})
        fill = ln.makeelement(qn("a:solidFill"), {})
        fill.append(fill.makeelement(qn("a:srgbClr"), {"val": t["ink"]}))
        ln.append(fill)
        tcpr.insert(0, ln)


def _png_size(png: bytes) -> tuple[int, int]:
    return int.from_bytes(png[16:20], "big"), int.from_bytes(png[20:24], "big")


def _ppt_figure(slide, t: dict, fig: Figure, label: str) -> None:
    from pptx.util import Inches

    w_px, h_px = _png_size(fig.png)
    box_w, box_h = _SLIDE_W - 1.2, _SLIDE_H - 2.75
    scale = min(box_w / w_px, box_h / h_px)
    w, h = w_px * scale, h_px * scale
    slide.shapes.add_picture(io.BytesIO(fig.png), Inches((_SLIDE_W - w) / 2), Inches(1.65), Inches(w), Inches(h))
    _ppt_text(slide, t, 0.6, 1.75 + h, _SLIDE_W - 1.2, 0.8, [f"{label}　{fig.caption}"], size=11, color="dim", markup=False)


def to_pptx(rep: Report, theme: str = DEFAULT_PPT_THEME) -> bytes:
    """簡報版報告：封面、重點摘要、各章節解讀、每張統計表 (過長自動分頁) 與 seaborn 圖、統計方法。"""
    from pptx import Presentation
    from pptx.util import Inches

    t = PPT_THEMES.get(theme, PPT_THEMES[DEFAULT_PPT_THEME])
    L = _labeler(rep)
    cont = lambda n: L(" (續 {n})", n=n + 1) if n else ""  # noqa: E731
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(_SLIDE_W), Inches(_SLIDE_H)
    _ppt_cover(prs, t, rep)

    for i, chunk in enumerate(_chunks(rep.highlights)):
        s = _ppt_slide(prs, t, "Highlights", L("重點摘要") + cont(i))
        _ppt_text(s, t, 0.6, 1.7, 12.1, 5.3, [f"•  {h}" for h in chunk], size=t["body_size"] + 1, gap=10)

    for i, sec in enumerate(rep.sections, 1):
        kicker = f"{i:02d}"
        if t["dividers"]:
            _ppt_divider(prs, t, kicker, sec.title)
        paras = list(sec.paragraphs) + [(f"⚠ {w}", "warn") for w in sec.warnings]
        for k, chunk in enumerate(_chunks(paras)):
            s = _ppt_slide(prs, t, kicker, sec.title + cont(k))
            _ppt_text(s, t, 0.6, 1.7, 12.1, 5.3, chunk, size=t["body_size"], gap=10)
        for j, (name, tab) in enumerate(sec.tables, 1):
            parts = [tab.iloc[a:a + _TABLE_ROWS] for a in range(0, max(len(tab), 1), _TABLE_ROWS)]
            for k, part in enumerate(parts):
                more = f"　({k + 1}/{len(parts)})" if len(parts) > 1 else ""
                s = _ppt_slide(prs, t, f"{kicker} · Table {i}.{j}", L("表 {no}", no=f"{i}.{j}") + f"　{name}{more}")
                _ppt_table(s, t, part)
        for j, fig in enumerate(sec.figures, 1):
            s = _ppt_slide(prs, t, f"{kicker} · Figure {i}.{j}", sec.title)
            _ppt_figure(s, t, fig, L("圖 {no}", no=f"{i}.{j}"))

    s = _ppt_slide(prs, t, f"{len(rep.sections) + 1:02d}", L("統計方法"))
    _ppt_text(s, t, 0.6, 1.7, 12.1, 4.9, rep.methods, size=15, gap=8)
    _ppt_text(s, t, 0.6, 6.55, 12.1, 0.5, [rep.software], size=10, color="dim", markup=False)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def to_excel(rep: Report) -> bytes:
    return data.to_excel_bytes({name: tab for name, tab in rep.all_tables()})
