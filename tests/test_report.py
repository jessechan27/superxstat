"""自動報告：檢定選擇邏輯、報告數值與單項分析一致、三種匯出格式。"""

import io

import matplotlib

matplotlib.use("Agg")

import numpy as np
import pandas as pd
import pytest
from scipy import stats

from biostat import data, report
from biostat import stats as bs

DEMO = data.make_demo_data()


@pytest.mark.parametrize("strategy,normal,homog_p,k,expected", [
    ("auto", True, 0.5, 2, "student"),
    ("auto", True, 0.01, 2, "welch"),
    ("auto", False, 0.5, 2, "mwu"),
    ("auto", None, 0.5, 2, "student"),  # 樣本太少無法檢定常態 → 不推翻母數假設
    ("auto", True, 0.5, 4, "anova"),
    ("auto", True, 0.001, 4, "welch_anova"),
    ("auto", False, 0.001, 4, "kw"),
    ("parametric", False, 0.5, 3, "anova"),
    ("nonparametric", True, 0.5, 2, "mwu"),
])
def test_choose(strategy, normal, homog_p, k, expected):
    assert report._choose(strategy, normal, homog_p, k) == expected


def _section(rep, prefix):
    return next(s for s in rep.sections if s.title.startswith(prefix))


def test_groups_report_matches_single_analyses():
    df = DEMO["Groups"]
    rep = report.build_report(df, report.ReportOptions(["TumorVolume", "IL6"], group="Group", categorical=["Sex"],
                                                       figures=False), "Groups")
    comp = dict(_section(rep, "組間比較").tables)["組間比較結果"].set_index("變項")
    xs = {v: [df.loc[df.Group == g, v].dropna().to_numpy() for g in plots_levels(df)] for v in ("TumorVolume", "IL6")}
    for v, row in comp.iterrows():
        ref = stats.kruskal(*xs[v]).pvalue if row["檢定方法"].startswith("Kruskal") else stats.f_oneway(*xs[v]).pvalue
        assert row["p 值"] == pytest.approx(ref)
    assert {s.title.split("：")[0] for s in rep.sections} >= {"資料概覽", "描述性統計", "前提檢定", "組間比較", "相關分析 (Spearman)",
                                                            "類別變項分析"}
    assert rep.highlights and rep.methods
    # 事後檢定只在整體檢定顯著時列出
    names = [n for n, _ in _section(rep, "組間比較").tables]
    for v, row in comp.iterrows():
        assert any(n.startswith(f"{v} 事後檢定") for n in names) == (row["p 值"] < 0.05)


def test_nonsignificant_outcome_has_no_posthoc():
    df = DEMO["Groups"]
    rep = report.build_report(df, report.ReportOptions(["Weight_g"], group="Group", figures=False), "Groups")
    sec = _section(rep, "組間比較")
    assert dict(sec.tables)["組間比較結果"]["結論"].iloc[0] == "不顯著"
    assert not any("事後檢定" in n for n, _ in sec.tables)


def plots_levels(df):
    from biostat import plots

    return plots.category_levels(df["Group"])


def test_two_group_report_uses_welch_when_variances_differ():
    rng = np.random.default_rng(7)
    df = pd.DataFrame({"g": ["A"] * 30 + ["B"] * 30, "y": np.r_[rng.normal(0, 1, 30), rng.normal(1, 4, 30)]})
    rep = report.build_report(df, report.ReportOptions(["y"], group="g", figures=False, correlation=False))
    row = dict(_section(rep, "組間比較").tables)["組間比較結果"].iloc[0]
    a, b = df.y[:30], df.y[30:]
    if stats.levene(a, b, center="median").pvalue < 0.05 and min(stats.shapiro(a).pvalue, stats.shapiro(b).pvalue) >= 0.05:
        assert row["檢定方法"] == "Welch t 檢定"
        assert row["p 值"] == pytest.approx(stats.ttest_ind(a, b, equal_var=False).pvalue)


def test_numeric_group_codes_and_fisher():
    # 期望次數 Y 欄 = 8 × 7 / 16 = 3.5 < 5 → 應改用 Fisher 精確檢定
    df = pd.DataFrame({"trt": [0] * 8 + [1] * 8, "resp": ["Y"] * 1 + ["N"] * 7 + ["Y"] * 6 + ["N"] * 2,
                       "y": np.arange(16, dtype=float)})
    rep = report.build_report(df, report.ReportOptions(["y"], group="trt", categorical=["resp"], figures=False))
    cat = dict(_section(rep, "類別變項分析").tables)["類別變項關聯摘要"].iloc[0]
    ct = pd.crosstab(df.trt, df.resp).to_numpy()
    assert cat["檢定"] == "Fisher 精確檢定"
    assert cat["p 值"] == pytest.approx(stats.fisher_exact(ct)[1])


def test_report_without_group():
    rep = report.build_report(DEMO["Groups"], report.ReportOptions(["IL6", "TNFa"], figures=False))
    titles = [s.title for s in rep.sections]
    assert not any(t.startswith("組間比較") for t in titles) and any(t.startswith("相關分析") for t in titles)


def test_invalid_options():
    with pytest.raises(ValueError):
        report.build_report(DEMO["Groups"], report.ReportOptions([]))
    with pytest.raises(ValueError):
        report.build_report(DEMO["Groups"], report.ReportOptions(["IL6"], group="SubjectID"))


def test_exports_with_figures():
    from docx import Document
    from openpyxl import load_workbook

    rep = report.build_report(DEMO["Groups"], report.ReportOptions(["TumorVolume", "IL6"], group="Group", categorical=["Sex"]))
    assert sum(len(s.figures) for s in rep.sections) >= 6
    assert all(f.png[:8] == b"\x89PNG\r\n\x1a\n" for s in rep.sections for f in s.figures)
    page = report.to_html(rep)
    assert page.count("<figure>") == sum(len(s.figures) for s in rep.sections) and "統計方法" in page
    doc = Document(io.BytesIO(report.to_docx(rep)))
    assert len(doc.inline_shapes) == sum(len(s.figures) for s in rep.sections)
    assert len(doc.tables) == len(rep.all_tables())
    wb = load_workbook(io.BytesIO(report.to_excel(rep)))
    assert len(wb.sheetnames) == len(rep.all_tables())


@pytest.mark.parametrize("theme", list(report.PPT_THEMES))
def test_pptx_export(theme):
    from pptx import Presentation

    rep = report.build_report(DEMO["Groups"], report.ReportOptions(["TumorVolume", "IL6"], group="Group", categorical=["Sex"]))
    prs = Presentation(io.BytesIO(report.to_pptx(rep, theme)))
    pics = sum(1 for s in prs.slides for sh in s.shapes if sh.shape_type == 13)  # 13 = PICTURE
    tables = [sh.table for s in prs.slides for sh in s.shapes if sh.has_table]
    assert pics == sum(len(s.figures) for s in rep.sections)
    # 長表格分頁後，資料列總數仍與原表相同
    assert sum(len(t.rows) - 1 for t in tables) == sum(max(len(tab), 1) for _, tab in rep.all_tables())
    assert all(len(t.rows) - 1 <= report._TABLE_ROWS for t in tables)
    texts = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert rep.title in texts and "統計方法" in texts and "重點摘要" in texts
