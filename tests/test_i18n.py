"""中英文切換：對照表樣板一致、英文報告沒有殘留中文、各頁面在英文模式可正常執行。"""

import io
import re
import string

import matplotlib

matplotlib.use("Agg")

import pytest
from streamlit.testing.v1 import AppTest

from biostat import data, i18n, report
from biostat.locales.en import EN

CJK = re.compile(r"[　-〿一-鿿＀-￯]")
DEMO = data.make_demo_data()


@pytest.fixture
def english():
    i18n.set_lang("en")
    yield
    i18n.set_lang("zh")


def _fields(s: str) -> set:
    return {f for _a, f, _b, _c in string.Formatter().parse(s) if f}


def test_catalog_placeholders_match():
    assert [k for k, v in EN.items() if _fields(k) != _fields(v)] == []


# 已完成英文化的檔案：程式中每個 _() / _t() / L() 的中文字串都必須在對照表裡
TRANSLATED_FILES = ["app.py", "biostat/report.py", "biostat/stats.py", "biostat/data.py", "biostat/views/common.py",
                    "biostat/views/page_data.py", "biostat/views/page_report.py", "biostat/views/page_stats.py"]


def test_every_translated_string_has_english():
    import ast
    import pathlib

    from biostat.views import page_data, page_stats

    root = pathlib.Path(__file__).resolve().parents[1]
    keys = list(page_stats.ANALYSES) + list(page_data.KINDS) + list(report.STRATEGIES) + list(report.PPT_THEMES)
    for f in TRANSLATED_FILES:
        for n in ast.walk(ast.parse((root / f).read_text(encoding="utf-8"))):
            if isinstance(n, ast.Call) and isinstance(n.func, ast.Name) and n.func.id in ("_", "_t", "L") and n.args:
                a = n.args[0]
                options = [a] if isinstance(a, ast.Constant) else [a.body, a.orelse] if isinstance(a, ast.IfExp) else \
                    list(a.value.values) if isinstance(a, ast.Subscript) and isinstance(a.value, ast.Dict) else []
                keys += [o.value for o in options if isinstance(o, ast.Constant) and isinstance(o.value, str)]
    missing = sorted({k for k in keys if CJK.search(k) and k not in EN})
    assert missing == []


def test_translate_and_fallback(english):
    assert i18n._("重點摘要") == "Key findings"
    assert i18n._("表 {no}", no="1.2") == "Table 1.2"
    assert i18n._("尚未翻譯的字串") == "尚未翻譯的字串"  # 查不到就保持原文


def _report_texts(rep) -> list[str]:
    texts = [rep.title] + rep.highlights + rep.methods + [rep.software]
    for s in rep.sections:
        texts += [s.title] + s.paragraphs + s.warnings + [f.caption for f in s.figures]
        for name, tab in s.tables:
            texts += [name] + [str(c) for c in tab.columns]
            texts += [v for c in tab.columns if tab[c].dtype == object for v in tab[c] if isinstance(v, str)]
    texts.append(re.sub(r'src="data:image/png;base64,[^"]*"', "", report.to_html(rep)))
    return texts


@pytest.mark.parametrize("sheet,y,g,cats", [
    ("Groups", ["TumorVolume", "IL6", "TNFa"], "Group", ["Sex"]),
    ("Clinical", ["Age"], "Treatment", ["Sex", "Smoker", "Outcome"]),
    ("Paired", ["Before", "After"], None, []),
])
def test_english_report_has_no_chinese(english, sheet, y, g, cats):
    rep = report.build_report(DEMO[sheet], report.ReportOptions(y, group=g, categorical=cats, figures=False), sheet)
    assert rep.lang == "en"
    leaks = sorted({w for t in _report_texts(rep) for w in t.split() if CJK.search(w)})
    assert leaks == []


def test_exports_keep_report_language():
    """報告用英文產生後，即使介面切回中文，匯出的標籤仍是英文。"""
    i18n.set_lang("en")
    rep = report.build_report(DEMO["Groups"], report.ReportOptions(["IL6"], group="Group", figures=False))
    i18n.set_lang("zh")
    page = report.to_html(rep)
    assert "Key findings" in page and "重點摘要" not in page and 'lang="en"' in page
    from docx import Document

    doc = Document(io.BytesIO(report.to_docx(rep)))
    assert "Key findings" in [p.text for p in doc.paragraphs]


SCRIPT = """
import matplotlib
matplotlib.use("Agg")
import streamlit as st
from biostat import i18n
from biostat.views import common as ui, {page}
i18n.set_lang(st.session_state.get("lang", "zh"))
ui.language_switch()
ds = ui.sidebar_data()
{page}.render(ds)
"""


@pytest.mark.parametrize("page", ["page_data", "page_report", "page_stats"])
def test_pages_render_in_english(page):
    # 註：AppTest 無法取得 app 的 format_func，在英文模式點按鈕會誤判選單值；
    # 英文報告的產生流程已由 test_english_report_has_no_chinese 直接測試。
    at = AppTest.from_string(SCRIPT.format(page=page), default_timeout=120)
    at.session_state["lang"] = "en"
    at.run()
    assert not at.exception, [e.value for e in at.exception]
    shown = " ".join(m.value for m in at.markdown)
    assert "Data source" in shown and "資料來源" not in shown
